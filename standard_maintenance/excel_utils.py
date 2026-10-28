# -*- coding: utf-8 -*-
"""Excel 读取/写入工具（标准维护）。

读取：用 openpyxl 解析输入 Excel（第1行合并表头、第2行列名、第3行起数据），
按精确列名匹配目标列，不依赖 pandas（规避 numpy/pandas 二进制兼容问题）。
写入：完整复制输入文件，末尾追加"变更标准"组 6 列（第1行合并组表头）。
"""

import shutil

from openpyxl import load_workbook

from standard_maintenance.constants import (
    INPUT_COLUMNS,
    RESULT_COLUMNS,
    GROUP_CHANGE_STANDARD,
    COL_STD_ID,
)
from standard_maintenance.state import DictEntryItem
from common.llm_client import append_module_log, format_token_stats_summary


def _safe_str(val) -> str:
    """安全转字符串，处理 None 和空值。"""
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ("nan", "none", "nat"):
        return ""
    return s


def _match_columns(header: list) -> dict[str, str]:
    """按精确列名匹配 Excel 列名到内部字段名。

    Returns:
        {"field_name": "实际列名", ...}
    """
    actual_cols = set(str(c) for c in header if c is not None)
    col_map = {}
    for internal_name, keywords in INPUT_COLUMNS.items():
        for kw in keywords:
            if kw in actual_cols:
                col_map[internal_name] = kw
                break
    return col_map


def read_excel(file_path: str) -> list[DictEntryItem]:
    """读取 Excel（openpyxl），返回 DictEntryItem 列表。

    输入 Excel 第1行为合并表头，第2行为列名，从第3行开始为数据。
    """
    wb = load_workbook(file_path, read_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter)  # 第1行：合并组表头（忽略）
    col_name_row = next(rows_iter)  # 第2行：列名
    col_map = _match_columns(col_name_row)

    # 校验必填列是否全部匹配到
    required_keys = ["field_name", "field_type", "is_enum", "business_meaning"]
    missing = [k for k in required_keys if k not in col_map]
    if missing:
        wb.close()
        raise ValueError(
            f"输入Excel缺少必填列: {missing}，"
            f"请检查输入文件是否为质检输出格式（需含'中文字段名(必填)'、"
            f"'字段所属类型(必填)'、'是否枚举(必填)'、'业务定义(必填)'列）。"
        )

    col_idx = {name: i for i, name in enumerate(col_name_row)}

    rows: list[DictEntryItem] = []
    idx = 0
    for data_row in rows_iter:
        rows.append(DictEntryItem(
            index=idx,
            field_name=_safe_str(data_row[col_idx[col_map["field_name"]]]),
            field_type=_safe_str(data_row[col_idx[col_map["field_type"]]]),
            is_enum=_safe_str(data_row[col_idx[col_map["is_enum"]]]),
            enum_values=_safe_str(data_row[col_idx[col_map["enum_values"]]]) if "enum_values" in col_map else "",
            business_meaning=_safe_str(data_row[col_idx[col_map["business_meaning"]]]),
            # 初始化结果字段
            normalized_name="",
            generated_definition="",
            final_definition="",
        ))
        idx += 1

    wb.close()

    print(f"已读取 {len(rows)} 行数据")
    print(f"  列映射: {col_map}")
    return rows


def write_excel(
    file_path: str,
    rows: list[DictEntryItem],
    input_file: str,
    append: bool = False,
):
    """将"变更标准"结果写入 Excel。

    完整复制输入文件作为输出（保留全部原始列、数据、合并表头和格式），
    在末尾追加"变更标准"组 6 列：
      标准编号（留空）/ 中文字段名（规范化后）/ 字段所属类型 / 是否枚举 / 枚举值 / 业务定义（最终整合）。
    第1行对 6 列合并并写组名"变更标准"，第2行写各列名。

    append=True 时（增量写盘）：不重建文件，打开已存在的输出文件，
    定位"标准编号"列并覆盖填充当前 rows 的结果。用于分批处理时每完成一批即写入一次，
    重复写入已处理行是幂等的。
    """
    result_cols = RESULT_COLUMNS

    if append:
        # 增量写盘：打开已有输出文件，定位"变更标准"组起始列（第2行"标准编号"列名）
        wb = load_workbook(file_path)
        ws = wb.active
        col_offset = None
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=2, column=col).value == COL_STD_ID:
                col_offset = col
                break
        if col_offset is None:
            raise ValueError(
                f"增量写入失败：输出文件 {file_path} 未找到结果列 '{COL_STD_ID}'"
            )
        # 补齐缺失的结果列名（增量时结果列定义与首次一致，防御性补齐）
        for i, col_name in enumerate(result_cols):
            if ws.cell(row=2, column=col_offset + i).value in (None, ""):
                ws.cell(row=2, column=col_offset + i, value=col_name)
    else:
        # 1. 完整复制输入文件作为输出文件
        shutil.copyfile(input_file, file_path)

        # 2. 打开输出文件，追加"变更标准"列
        wb = load_workbook(file_path)
        ws = wb.active
        max_col = ws.max_column

        # 3. 第2行追加列名
        col_offset = max_col + 1
        for i, col_name in enumerate(result_cols):
            ws.cell(row=2, column=col_offset + i, value=col_name)

        # 4. 第1行追加"变更标准"合并组表头
        start_col = col_offset
        end_col = col_offset + len(result_cols) - 1
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        ws.cell(row=1, column=start_col, value=GROUP_CHANGE_STANDARD)

    # 5. 从第3行起按行 index 填入结果
    # （增量写盘时 rows 可能只含当前批次行，须按 index 定位绝对行号，不能按列表顺序）
    for r in sorted(rows, key=lambda r: r["index"]):
        excel_row = 3 + r["index"]
        ws.cell(row=excel_row, column=col_offset, value="")          # 标准编号留空
        ws.cell(row=excel_row, column=col_offset + 1, value=r["normalized_name"])
        ws.cell(row=excel_row, column=col_offset + 2, value=r["field_type"])
        ws.cell(row=excel_row, column=col_offset + 3, value=r["is_enum"])
        ws.cell(row=excel_row, column=col_offset + 4, value=r["enum_values"])
        ws.cell(row=excel_row, column=col_offset + 5, value=r["final_definition"])

    wb.save(file_path)
    print(f"结果已写入: {file_path}")

    # 打印统计
    sorted_rows = sorted(rows, key=lambda r: r["index"])
    total = len(sorted_rows)
    normalized = sum(1 for r in sorted_rows if r["normalized_name"] != r["field_name"])
    merged = sum(1 for r in sorted_rows if r["final_definition"] != r["generated_definition"])
    print(f"  总计: {total} 行 | 字段名规范化改写: {normalized} | 业务定义补充整合: {merged}")

    # 每次写盘时，把当前累计的 token 统计写入模块日志（命中/未命中缓存输入、输出 token）
    try:
        token_log = format_token_stats_summary()
        print(token_log)
        append_module_log("standard_maintenance", "\n" + token_log + "\n")
    except Exception as e:
        print(f"  [WARNING] 写入 token 统计日志失败: {e}")
